"""
AccessTrap-CS — Corpus Generation Script
Generates responses: concepts x 3 conditions x 7 models (35 concepts -> 735).
Requires: Ollama running on localhost:11434, packages: requests, openpyxl
Usage: python generate.py
"""

import csv
import os
import re
import sys
import time
import requests
from datetime import datetime

# ── CONFIG ────────────────────────────────────────────────────────────────────

OLLAMA_CHAT_URL = "http://localhost:11434/api/chat"
OLLAMA_TAGS_URL = "http://localhost:11434/api/tags"
OLLAMA_GEN_URL  = "http://localhost:11434/api/generate"

MODELS = [
    "mistral:7b",
    "deepseek-r1:8b",
    "ministral-3:14b",
    "phi4:14b",
    "gpt-oss:20b",
    "gemma4:31b",
    "deepseek-r1:32b",
]

PROMPTS = {
    "A_standard": (
        "Explain the following computer science concept accurately and clearly, "
        "suitable for a first-year computer science student.\n\n"
        "Answer in English.\n\n"
        "Concept: {concept}"
    ),
    "B_simplified": (
        "Explain the following computer science concept to a student who has "
        "reading difficulties. Use simplified language, short sentences, and "
        "concrete analogies instead of code or formal definitions. Avoid "
        "technical jargon where possible.\n\n"
        "Answer in English.\n\n"
        "Concept: {concept}"
    ),
    "C_highly_simplified": (
        "Explain the following computer science concept to a student with "
        "significant attention difficulties. Keep the explanation within five "
        "sentences. Use one single, vivid analogy that captures the core idea. "
        "Make it memorable and engaging.\n\n"
        "Answer in English.\n\n"
        "Concept: {concept}"
    ),
}

INPUT_FILE = "ground_truth.xlsx"
OUTPUT_CSV = "corpus.csv"
LOG_FILE   = "generate_log.txt"

TEMPERATURE = 0.7
SEED        = 42        # retries use SEED+attempt-1 so a retry is not identical
NUM_PREDICT = -1        # no token limit
NUM_CTX     = 16384     # must fit prompt + thinking + answer; lower if memory-bound
KEEP_ALIVE  = "30m"
TIMEOUT     = 1800      # max wait per request (safety net against runaway output)
MAX_ATTEMPTS       = 3
RETRY_DELAY        = 5
MIN_RESPONSE_CHARS = 20

CSV_FIELDS = [
    "concept_id", "domain", "concept", "condition", "model",
    "prompt", "response",
    "done_reason", "eval_count", "word_count", "truncated",
    "ground_truth", "reference", "timestamp",
]

# ── HELPERS ───────────────────────────────────────────────────────────────────

def log(msg):
    line = f"[{datetime.now().strftime('%H:%M:%S')}] {msg}"
    print(line, flush=True)
    with open(LOG_FILE, "a", encoding="utf-8") as f:
        f.write(line + "\n")


def clean_cell(value):
    return "" if value is None else str(value).strip()


def clean_gt(text):
    """Remove stray quote characters left around the ground truth text."""
    return text.strip().strip("'\"\u2018\u2019\u201c\u201d").strip()


def load_concepts(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    raw_headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]

    aliases = {"id": "id", "domain": "domain", "concept": "concept",
               "ground_truth": "ground_truth", "reference": "reference",
               "refrence": "reference"}
    headers = []
    for h in raw_headers:
        if h is None:
            headers.append(None)
        else:
            key = str(h).strip().lower().replace(" ", "_")
            headers.append(aliases.get(key, key))

    required = {"id", "domain", "concept", "ground_truth", "reference"}
    if not required.issubset(headers):
        log(f"FATAL: missing columns. Found {headers}, need {sorted(required)}")
        sys.exit(1)

    concepts = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        r = dict(zip(headers, row))
        if not clean_cell(r.get("id")) or not clean_cell(r.get("concept")):
            continue  # blank trailing rows
        concepts.append({
            "id":           clean_cell(r["id"]),
            "domain":       clean_cell(r["domain"]),
            "concept":      clean_cell(r["concept"]),
            "ground_truth": clean_gt(clean_cell(r["ground_truth"])),
            "reference":    clean_cell(r["reference"]),
        })

    ids = [c["id"] for c in concepts]
    dupes = {i for i in ids if ids.count(i) > 1}
    if dupes:
        log(f"FATAL: duplicate concept ids: {sorted(dupes)}")
        sys.exit(1)
    for c in concepts:
        if not c["ground_truth"]:
            log(f"WARNING: empty ground truth for {c['id']}")
    if not concepts:
        log("FATAL: no concepts loaded.")
        sys.exit(1)
    return concepts


def get_installed_models():
    try:
        r = requests.get(OLLAMA_TAGS_URL, timeout=10)
        r.raise_for_status()
    except Exception as e:
        log(f"FATAL: cannot reach Ollama: {e}")
        sys.exit(1)
    return {m["name"]: m.get("digest", "") for m in r.json().get("models", [])}


def verify_models(installed):
    missing = [m for m in MODELS if m not in installed]
    if missing:
        log("FATAL: these models are not installed (ollama pull <name>):")
        for m in missing:
            log(f"   - {m}")
        sys.exit(1)
    for m in MODELS:
        log(f"  {m:20s} digest {installed[m][:12]}")


def check_existing_csv():
    """Refuse to append to a CSV whose columns differ from CSV_FIELDS."""
    if not os.path.exists(OUTPUT_CSV):
        return
    with open(OUTPUT_CSV, newline="", encoding="utf-8-sig") as f:
        header = next(csv.reader(f), [])
    if header != CSV_FIELDS:
        log(f"FATAL: {OUTPUT_CSV} has a different header than this script.")
        log("       Rename or delete the old file and run again.")
        sys.exit(1)


def load_done():
    done = set()
    if not os.path.exists(OUTPUT_CSV):
        return done
    with open(OUTPUT_CSV, newline="", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            if len((row.get("response") or "").strip()) >= MIN_RESPONSE_CHARS:
                done.add((row["concept_id"], row["condition"], row["model"]))
    return done


THINK_BLOCK  = re.compile(r"<think>.*?</think>", re.DOTALL | re.IGNORECASE)
THINK_UNOPEN = re.compile(r"^.*?</think>", re.DOTALL | re.IGNORECASE)
THINK_UNCLOSED = re.compile(r"<think>.*$", re.DOTALL | re.IGNORECASE)


def strip_thinking(text):
    text = THINK_BLOCK.sub("", text)
    text = THINK_UNOPEN.sub("", text)      # closing tag without opening tag
    text = THINK_UNCLOSED.sub("", text)    # thinking never finished
    return text.strip()


def query_ollama(model, prompt, seed):
    payload = {
        "model": model,
        "messages": [{"role": "user", "content": prompt}],
        "stream": False,
        "keep_alive": KEEP_ALIVE,
        "options": {
            "temperature": TEMPERATURE,
            "seed": seed,
            "num_predict": NUM_PREDICT,
            "num_ctx": NUM_CTX,
        },
    }
    resp = requests.post(OLLAMA_CHAT_URL, json=payload, timeout=TIMEOUT)
    resp.raise_for_status()
    data = resp.json()
    # Newer Ollama: reasoning is in message.thinking (ignored here).
    # Older Ollama: <think>...</think> is inline in content (stripped here).
    content = strip_thinking((data.get("message") or {}).get("content") or "")
    done_reason = data.get("done_reason", "")
    eval_count = data.get("eval_count", 0)
    total_tokens = data.get("prompt_eval_count", 0) + eval_count
    # Context-limit hits are not always reported as done_reason == "length",
    # so also flag runs that filled the whole context window.
    truncated = done_reason == "length" or total_tokens >= NUM_CTX - 16
    return {"content": content, "done_reason": done_reason,
            "eval_count": eval_count, "truncated": truncated}


def generate_with_retry(model, prompt):
    last = ""
    for attempt in range(1, MAX_ATTEMPTS + 1):
        try:
            result = query_ollama(model, prompt, SEED + attempt - 1)
            if len(result["content"]) >= MIN_RESPONSE_CHARS:
                return result
            last = f"empty/short answer ({len(result['content'])} chars)"
        except requests.exceptions.ReadTimeout:
            raise RuntimeError(f"timed out after {TIMEOUT}s")  # don't retry
        except Exception as e:
            last = str(e)
        log(f"    attempt {attempt}/{MAX_ATTEMPTS} failed: {last}")
        time.sleep(RETRY_DELAY)
    raise RuntimeError(f"failed after {MAX_ATTEMPTS} attempts: {last}")


def unload_model(model):
    try:
        requests.post(OLLAMA_GEN_URL, json={"model": model, "keep_alive": 0}, timeout=60)
    except Exception:
        pass


def append_row(row):
    new_file = not os.path.exists(OUTPUT_CSV)
    # BOM only on file creation so Excel shows č, ć, š, ž, Θ correctly.
    with open(OUTPUT_CSV, "a", newline="",
              encoding="utf-8-sig" if new_file else "utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=CSV_FIELDS)
        if new_file:
            writer.writeheader()
        writer.writerow(row)

# ── MAIN ──────────────────────────────────────────────────────────────────────

def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    if not os.path.exists(INPUT_FILE):
        log(f"FATAL: {INPUT_FILE} not found.")
        sys.exit(1)

    log("Checking models...")
    verify_models(get_installed_models())
    check_existing_csv()

    concepts = load_concepts(INPUT_FILE)
    done = load_done()
    total = len(concepts) * len(PROMPTS) * len(MODELS)
    log("=" * 60)
    log(f"Plan: {len(concepts)} concepts x {len(PROMPTS)} conditions x "
        f"{len(MODELS)} models = {total} responses")
    log(f"Already done: {len(done)}")
    log("=" * 60)

    generated = skipped = errors = truncated = 0

    for model in MODELS:                       # one model at a time
        log(f"MODEL: {model}")
        for concept in concepts:
            for condition, template in PROMPTS.items():
                key = (concept["id"], condition, model)
                if key in done:
                    skipped += 1
                    continue

                prompt = template.format(concept=concept["concept"])
                log(f"  {concept['id']} | {condition} | {model}")
                try:
                    result = generate_with_retry(model, prompt)
                except Exception as e:
                    log(f"    FAILED: {e}")
                    errors += 1
                    continue

                if result["truncated"]:
                    truncated += 1
                    log("    WARNING: output may be cut off (context limit)")

                append_row({
                    "concept_id":   concept["id"],
                    "domain":       concept["domain"],
                    "concept":      concept["concept"],
                    "condition":    condition,
                    "model":        model,
                    "prompt":       prompt,
                    "response":     result["content"],
                    "done_reason":  result["done_reason"],
                    "eval_count":   result["eval_count"],
                    "word_count":   len(result["content"].split()),
                    "truncated":    result["truncated"],
                    "ground_truth": concept["ground_truth"],
                    "reference":    concept["reference"],
                    "timestamp":    datetime.now().isoformat(),
                })
                generated += 1
        unload_model(model)

    log("=" * 60)
    log(f"Done. Generated: {generated} | Skipped (already done): {skipped} | "
        f"Failed: {errors} | Possibly truncated: {truncated}")
    log(f"Output: {OUTPUT_CSV}")


if __name__ == "__main__":
    main()