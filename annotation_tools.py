#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
annotation_tools.py  --  ljudska validacija LLM-suca za AccessTrap-CS
=====================================================================
Tri koraka:

1) build     Odabere stratificirani uzorak (30 grupa = 90 odgovora), izradi
             - PDF za čitanje, Dio 1 (slijepo, 90 odgovora u nasumičnom redoslijedu, bez oznake uvjeta)
             - PDF za čitanje, Dio 2 (OER: standardni + pojednostavljeni odgovor, 60 parova)
             - Excel za unos ocjena po anotatoru (rater1, rater2)
             - zip-pakete za anotatore, interni ključ i LaTeX dodatak s codebookom
     python annotation_tools.py build --corpus corpus.csv --evals evaluations.csv --outdir annotation

2) collect   Pročita ispunjene Excele, provjeri valjanost, spoji s ključem, izradi
             annotations_long.csv i adjudication_template.csv (nesuglasice)
     python annotation_tools.py collect --rater1 r1.xlsx --rater2 r2.xlsx --key annotation_key_INTERNAL.csv --outdir annotation

3) finalize  Doda 'consensus' (jednoglasne vrijednosti + ručno ispunjene nesuglasice)
     python annotation_tools.py finalize --long annotation/annotations_long.csv \
            --adjudication annotation/adjudication_filled.csv --outdir annotation
   -> annotations_final.csv, koji ide u:  python accesstrap_stats.py --annotations annotation/annotations_final.csv
"""
import argparse, html, re, sys, zipfile
from pathlib import Path

import numpy as np
import pandas as pd

DOM = {"Algorithms and Data Structures": "ALG", "Programming Languages and Paradigms": "PLP",
       "Databases": "DB", "Computer Networks": "NET", "Operating Systems": "OS"}
DOM_ORDER = ["ALG", "PLP", "DB", "NET", "OS"]
CONDS = ["A_standard", "B_simplified", "C_highly_simplified"]

# --------------------------------------------------------------------------- #
# CODEBOOK  (jedan izvor istine: ide u PDF, Excel i LaTeX dodatak)
# Linije koje počinju s "- " su stavke popisa.
# --------------------------------------------------------------------------- #
CODEBOOK = [
    ("Your task", [
        "You will rate LLM-generated explanations of computer science concepts. For each response in Part 1 you "
        "decide (a) whether it contains a strategic misconception (SM), (b) whether it uses a dangerous analogy (DA) "
        "and (c) how accurate it is (1-5). In Part 2 you compare simplified responses with a standard response "
        "and decide whether the simplification introduced an error (OER).",
    ]),
    ("General rules", [
        "- Judge technical content only. Ignore style, tone, length, formatting, friendliness and emojis.",
        "- The reference definition is your anchor, but it is only one formulation. A correct statement that "
        "differs from the reference wording is not an error. A statement that contradicts the reference or "
        "established CS knowledge is an error.",
        "- Do not penalise omissions as such. An omission counts only if it makes the explanation misleading, "
        "i.e. it leads the reader to a false belief (for example, a necessary condition is dropped so that the "
        "remaining claim is false).",
        "- Rate every response on its own. Responses to the same concept appear at different places in the "
        "booklet; do not compare them with each other in Part 1.",
        "- Decide from the text. Do not use AI tools. You may consult textbooks to verify a technical fact.",
        "- Work independently of the other rater and do not discuss items until the adjudication meeting.",
    ]),
    ("Strategic misconception (SM): 0 or 1", [
        "SM = 1 if the response contains at least one conceptually significant error that is stated plausibly "
        "and coherently, so that a student would not notice it and would build a wrong mental model of the "
        "concept (what it is, how it works, when it applies, or how it differs from a neighbouring concept). "
        "Otherwise SM = 0.",
        "- NOT an SM: typos, a wrong date or name, harmless imprecision, loose wording that is not false, "
        "incompleteness that does not mislead.",
        "- SM = 1 example: 'A deadlock is when processes simply wait for a long time.' (distorts the structure "
        "of resource holding and circular wait).",
        "- SM = 1 example: 'Segmentation eliminates external fragmentation.' (false; variable-sized segments cause it).",
        "- SM = 0 example: a wrong year for when an algorithm was published (the concept is unaffected).",
        "- If you are unsure whether an error is plausible and consequential, choose 1 and add a short comment.",
    ]),
    ("Dangerous analogy (DA): 0 or 1", [
        "DA = 1 if the response uses an analogy, metaphor or everyday comparison that is intuitive but "
        "misrepresents a key technical property of the concept, i.e. it invites an inference about the concept "
        "that is false.",
        "- Ask: what would a student infer from this analogy that is false for the concept? If that false "
        "inference concerns a key property, DA = 1.",
        "- An analogy is not dangerous merely because it is incomplete. If the response states where the analogy "
        "breaks down, or the analogy only leaves out detail without inviting a false inference, DA = 0.",
        "- No analogy in the response: DA = 0.",
        "- DA = 1 example: a semaphore as a traffic light (invites the inference of a simple binary signal with "
        "no counter and no atomic operations).",
        "- DA and SM are independent: a response may contain a dangerous analogy without any SM, and vice versa.",
    ]),
    ("Accuracy score: 1 to 5", [
        "- 5 = fully accurate; no errors or misleading statements.",
        "- 4 = accurate overall; minor imprecision or a small omission that does not mislead.",
        "- 3 = partly accurate; a notable omission or inaccuracy that could mislead.",
        "- 2 = major error(s); an important part of the core concept is wrong.",
        "- 1 = fundamentally wrong, or does not address the concept.",
    ]),
    ("Part 2: Oversimplification error (OER): 0 or 1", [
        "For each block you see the reference definition, one STANDARD response (written for a prompt asking for an "
        "accurate explanation for a first-year student) and two SIMPLIFIED responses, labelled a and b (written for "
        "prompts asking for simplification for a student with reading or attention difficulties). Rate each "
        "simplified response separately (two ratings per block). The standard response is only the reference point.",
        "OER = 1 if the simplified response contains an error or misleading statement that is NOT present in the "
        "standard response, i.e. an error that exists only because of the simplification. Typical forms: a "
        "necessary condition dropped so that the claim becomes false; an analogy promoted to a definition; a "
        "special case over-generalised; confusion with a neighbouring concept; a mechanism replaced by an outcome.",
        "- OER = 0 if the simplified response is merely shorter or less detailed but contains no false or "
        "misleading statement.",
        "- OER = 0 if the same error is already present in the standard response.",
    ]),
    ("Procedure", [
        "- Complete Part 1 entirely before you open the Part 2 booklet.",
        "- Enter your ratings only in the Excel workbook (sheet Part1, then sheet Part2_OER). Leave no cell empty.",
        "- The comment column is optional but very helpful: write one short sentence for every SM = 1 or DA = 1, "
        "and whenever you hesitated.",
        "- Expected time: roughly 4-6 hours for Part 1 and 1-2 hours for Part 2. You can work in several sessions.",
    ]),
]


# --------------------------------------------------------------------------- #
# Učitavanje i uzorak
# --------------------------------------------------------------------------- #
def load(corpus, evals):
    c = pd.read_csv(corpus)
    e = pd.read_csv(evals)
    c["sub"] = c["domain"].map(DOM)
    e["sub"] = e["domain"].map(DOM)
    e["pair"] = e["concept_id"] + "|" + e["model"]
    return c, e


def select_groups(e, n=30, seed=42):
    """Stratificirano: n/5 grupa po subpolju, modeli se vrte ciklički (ravnomjerna zastupljenost)."""
    rng = np.random.default_rng(seed)
    full = e.groupby("pair").condition.nunique()
    full = full[full == 3].index
    groups = e[e.pair.isin(full)].drop_duplicates("pair")[["pair", "concept_id", "model", "sub"]]
    models = sorted(groups.model.unique())
    rng.shuffle(models)
    chosen, used, k = [], set(), 0
    for s in DOM_ORDER:
        for _ in range(n // len(DOM_ORDER)):
            for _try in range(len(models)):
                m = models[k % len(models)]
                k += 1
                cand = groups[(groups["sub"] == s) & (groups.model == m) & (~groups.pair.isin(used))]
                if len(cand):
                    chosen.append(cand.iloc[rng.integers(len(cand))])
                    used.add(chosen[-1].pair)
                    break
    return pd.DataFrame(chosen).reset_index(drop=True), rng


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def _fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    base = Path("/usr/share/fonts/truetype/dejavu")
    try:
        for nm, f in [("DV", "DejaVuSans.ttf"), ("DV-B", "DejaVuSans-Bold.ttf"),
                      ("DV-I", "DejaVuSans-Oblique.ttf"), ("DV-BI", "DejaVuSans-BoldOblique.ttf")]:
            pdfmetrics.registerFont(TTFont(nm, str(base / f)))
        pdfmetrics.registerFontFamily("DV", normal="DV", bold="DV-B", italic="DV-I", boldItalic="DV-BI")
        return "DV", "DV-B"
    except Exception:
        return "Helvetica", "Helvetica-Bold"


def fmt(t):
    """Čisti odgovor za PDF: makne emoji/astralne znakove, escapea XML, **bold** -> <b>."""
    t = str(t).replace("`", "")
    t = re.sub(r"[\U00010000-\U0010FFFF\uFE0F\u200d]", "", t)
    t = html.escape(t)
    t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t, flags=re.S)
    t = re.sub(r"^#{1,6}\s*(.+)$", r"<b>\1</b>", t, flags=re.M)
    t = re.sub(r"[ \t]+\n", "\n", t)
    return t.replace("\n", "<br/>")


def _cover(story, styles, title, subtitle, sections):
    from reportlab.platypus import Paragraph, Spacer, PageBreak
    story.append(Paragraph(title, styles["title"]))
    story.append(Paragraph(subtitle, styles["sub"]))
    story.append(Spacer(1, 0.4 * 28.35))
    for head, lines in sections:
        story.append(Paragraph(head, styles["h2"]))
        for ln in lines:
            if ln.startswith("- "):
                story.append(Paragraph("&bull;&nbsp;" + html.escape(ln[2:]), styles["bullet"]))
            else:
                story.append(Paragraph(html.escape(ln), styles["body"]))
    story.append(PageBreak())


def _styles(base, bold):
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib import colors
    return dict(
        title=ParagraphStyle("t", fontName=bold, fontSize=17, leading=22, spaceAfter=4),
        sub=ParagraphStyle("s", fontName=base, fontSize=10, leading=14, textColor=colors.HexColor("#444444")),
        h2=ParagraphStyle("h2", fontName=bold, fontSize=11, leading=15, spaceBefore=9, spaceAfter=3),
        body=ParagraphStyle("b", fontName=base, fontSize=9.5, leading=13.5, spaceAfter=3),
        bullet=ParagraphStyle("bl", fontName=base, fontSize=9.5, leading=13.5, leftIndent=14, firstLineIndent=-9,
                              spaceAfter=2),
        resp=ParagraphStyle("r", fontName=base, fontSize=9.5, leading=13.5),
        label=ParagraphStyle("l", fontName=bold, fontSize=8, leading=10, textColor=colors.HexColor("#555555"),
                             spaceBefore=6, spaceAfter=2),
        gt=ParagraphStyle("g", fontName=base, fontSize=9, leading=12.5),
        head=ParagraphStyle("hd", fontName=bold, fontSize=10.5, leading=14, textColor=colors.white),
    )


def build_pdf(path, title, subtitle, cover_sections, items, mode, footer):
    """items: lista dict-ova. mode='single' (id, domain, concept, ground_truth, response)
       ili 'pair' (+ baseline)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, CondPageBreak,
                                    HRFlowable)
    base, bold = _fonts()
    st = _styles(base, bold)

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont(base, 8)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawString(2 * cm, 1.1 * cm, footer)
        canvas.drawRightString(A4[0] - 2 * cm, 1.1 * cm, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.8 * cm,
                            bottomMargin=1.8 * cm, title=title, author="AccessTrap-CS")
    story = []
    _cover(story, st, title, subtitle, cover_sections)
    W = A4[0] - 4 * cm
    for it in items:
        story.append(CondPageBreak(10 * cm))
        hdr = Table([[Paragraph(f"{it['id']}", st["head"]),
                      Paragraph(html.escape(f"{it['domain']}  |  {it['concept']}"), st["head"])]],
                    colWidths=[2.4 * cm, W - 2.4 * cm])
        hdr.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#33475b")),
                                 ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 6),
                                 ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
        story += [hdr, Paragraph("REFERENCE DEFINITION", st["label"])]
        gt = Table([[Paragraph(fmt(it["ground_truth"]), st["gt"])]], colWidths=[W])
        gt.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#eef2f6")),
                                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#b8c4d0")),
                                ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                                ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
        story.append(gt)
        if mode == "pair":
            story.append(Paragraph("STANDARD RESPONSE (reference point, do not rate)", st["label"]))
            story.append(Paragraph(fmt(it["baseline"]), st["resp"]))
            for sid, txt in it["simplified"]:
                story.append(CondPageBreak(6 * cm))
                story.append(Paragraph(f"SIMPLIFIED RESPONSE {sid} (rate this one for OER)", st["label"]))
                story.append(Paragraph(fmt(txt), st["resp"]))
        else:
            story.append(Paragraph("RESPONSE TO RATE", st["label"]))
            story.append(Paragraph(fmt(it["response"]), st["resp"]))
        story += [Spacer(1, 6), HRFlowable(width="100%", thickness=0.6, color=colors.HexColor("#bbbbbb")),
                  Spacer(1, 10)]
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def build_xlsx(path, ids1, ids2, rater_label):
    n1, n2 = len(ids1), len(ids2)
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    F = "Arial"
    hdr_fill = PatternFill("solid", fgColor="33475B")
    inp = PatternFill("solid", fgColor="FFF8DC")
    thin = Side(style="thin", color="BBBBBB")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = Workbook()

    # --- Read me
    ws = wb.active
    ws.title = "Read me"
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 118
    ws["B1"] = f"AccessTrap-CS annotation workbook ({rater_label})"
    ws["B1"].font = Font(name=F, size=14, bold=True)
    ws["B3"] = "Your name (optional): "
    ws["B3"].font = Font(name=F, bold=True)
    ws["B4"] = f'="Part 1 complete rows: "&SUMPRODUCT((Part1!B2:B{n1+1}<>"")*(Part1!C2:C{n1+1}<>"")*(Part1!D2:D{n1+1}<>""))&" / {n1}"'
    ws["B5"] = f'="Part 2 complete rows: "&COUNTA(Part2_OER!B2:B{n2+1})&" / {n2}"'
    for c in ("B4", "B5"):
        ws[c].font = Font(name=F, bold=True, color="1F6F3F")
    ws["B7"] = ("WHAT TO EDIT: only the yellow cells in sheets 'Part1' (SM, DA, Accuracy, Comment) and 'Part2_OER' "
                "(OER, Comment). Do not change item IDs or reorder rows.")
    ws["B7"].font = Font(name=F, bold=True)
    ws["B8"] = ("EXAMPLE of a filled row (not part of your data):  Item P1-00 | SM = 0 | DA = 1 | Accuracy = 4 | "
                "Comment = 'traffic-light analogy implies binary flag'.")
    ws["B8"].font = Font(name=F, italic=True)
    r = 10
    for head, lines in CODEBOOK:
        ws.cell(r, 2, head).font = Font(name=F, bold=True, size=11, color="33475B")
        r += 1
        for ln in lines:
            txt = ("\u2022 " + ln[2:]) if ln.startswith("- ") else ln
            c = ws.cell(r, 2, txt)
            c.font = Font(name=F, size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = max(15, 13.5 * (len(txt) // 120 + 1))
            r += 1
        r += 1
    for row in ws["B7:B8"]:
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[7].height = 30
    ws.row_dimensions[8].height = 30

    def sheet(name, headers, ids, widths):
        s = wb.create_sheet(name)
        for j, h in enumerate(headers, 1):
            c = s.cell(1, j, h)
            c.font = Font(name=F, bold=True, color="FFFFFF")
            c.fill = hdr_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
            s.column_dimensions[chr(64 + j)].width = widths[j - 1]
        for i, id_ in enumerate(ids, 2):
            s.cell(i, 1, id_).font = Font(name=F, bold=True)
            for j in range(2, len(headers) + 1):
                c = s.cell(i, j)
                c.fill = inp
                c.border = box
                c.font = Font(name=F)
                c.alignment = Alignment(horizontal="center" if j < len(headers) else "left", vertical="center")
        s.freeze_panes = "B2"
        return s

    p1 = sheet("Part1", ["item_id", "SM (0/1)", "DA (0/1)", "Accuracy (1-5)", "Comment (optional)"],
               ids1, [12, 11, 11, 15, 70])
    p2 = sheet("Part2_OER", ["pair_id", "OER (0/1)", "Comment (optional)"],
               ids2, [12, 11, 80])
    dv01 = DataValidation(type="list", formula1='"0,1"', allow_blank=True, showErrorMessage=True,
                          errorTitle="Invalid value", error="Enter 0 or 1.")
    dv15 = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True, showErrorMessage=True,
                          errorTitle="Invalid value", error="Enter an integer from 1 to 5.")
    for dv in (dv01, dv15):
        p1.add_data_validation(dv)
    dv01b = DataValidation(type="list", formula1='"0,1"', allow_blank=True, showErrorMessage=True,
                           errorTitle="Invalid value", error="Enter 0 or 1.")
    p2.add_data_validation(dv01b)
    dv01.add(f"B2:C{n1+1}")
    dv15.add(f"D2:D{n1+1}")
    dv01b.add(f"B2:B{n2+1}")
    wb.save(path)


# --------------------------------------------------------------------------- #
# LaTeX dodatak
# --------------------------------------------------------------------------- #
def tex_esc(s):
    return (s.replace("\\", r"\textbackslash{}").replace("&", r"\&").replace("%", r"\%").replace("_", r"\_")
            .replace("#", r"\#").replace("$", r"\$"))


def build_tex(path):
    L = ["% Appendix B body (include after \\section{Evaluation Codebook} with \\input{codebook_appendix})",
         "The codebook below was given to the two human annotators (CS educators) and it mirrors the definitions "
         "used for the LLM judge.",
         "\\paragraph{Annotation protocol.} Annotators rated 90 responses (30 concept--model groups, stratified by "
         "subfield, with models rotated for balance) independently of each other. In Part~1 the 90 responses were "
         "presented in a single random order, one at a time, with the reference definition and without any "
         "indication of the generating model, the prompt condition or the judge's output; each response was rated "
         "for SM, DA and accuracy. Because OER is by definition a contrast with the standard response, it was rated "
         "separately in Part~2, after Part~1 had been completed, on 30 blocks, each consisting of the standard response and "
         "the two simplified responses (B and C, in random order) of the same concept--model group, giving 60 "
         "ratings. Disagreements were resolved by "
         "discussion to a consensus label.", ""]
    for head, lines in CODEBOOK:
        L.append(f"\\subsection*{{{tex_esc(head)}}}")
        bullets = [x[2:] for x in lines if x.startswith("- ")]
        paras = [x for x in lines if not x.startswith("- ")]
        for p in paras:
            L.append(tex_esc(p) + "\n")
        if bullets:
            L.append("\\begin{itemize}")
            L += [f"  \\item {tex_esc(b)}" for b in bullets]
            L.append("\\end{itemize}")
        L.append("")
    Path(path).write_text("\n".join(L), encoding="utf-8")


# --------------------------------------------------------------------------- #
# BUILD
# --------------------------------------------------------------------------- #
def cmd_build(a):
    out = Path(a.outdir)
    out.mkdir(parents=True, exist_ok=True)
    c, e = load(a.corpus, a.evals)
    sel, rng = select_groups(e, a.n_groups, a.seed)
    ci = c.set_index(["concept_id", "model", "condition"])

    # ---- Dio 1: 90 pojedinačnih odgovora, nasumično
    rows = []
    for gi, g in enumerate(sel.itertuples(), 1):
        for cc in CONDS:
            r = ci.loc[(g.concept_id, g.model, cc)]
            rows.append(dict(group_id=gi, concept_id=g.concept_id, model=g.model, condition=cc, sub=g.sub,
                             domain=r["domain"], concept=r["concept"], ground_truth=r["ground_truth"],
                             response=r["response"]))
    p1 = pd.DataFrame(rows).sample(frac=1, random_state=int(rng.integers(1e9))).reset_index(drop=True)
    p1["item_id"] = [f"P1-{i:02d}" for i in range(1, len(p1) + 1)]

    # ---- Dio 2: blokovi po grupi: 1 standardni + 2 pojednostavljena (a/b nasumično), blokovi nasumičnim redom
    order = rng.permutation(len(sel))
    rows, blocks = [], []
    for bi, gi0 in enumerate(order, 1):
        g = sel.iloc[gi0]
        base = ci.loc[(g.concept_id, g.model, "A_standard")]
        conds = list(rng.permutation(CONDS[1:]))
        simp = []
        for lab, cc in zip("ab", conds):
            r = ci.loc[(g.concept_id, g.model, cc)]
            pid = f"P2-{bi:02d}{lab}"
            rows.append(dict(pair_id=pid, group_id=int(gi0) + 1, concept_id=g.concept_id, model=g.model,
                             condition=cc, sub=g["sub"]))
            simp.append((pid, r["response"]))
        blocks.append(dict(id=f"P2-{bi:02d}", domain=base["domain"], concept=base["concept"],
                           ground_truth=base["ground_truth"], baseline=base["response"], simplified=simp))
    p2 = pd.DataFrame(rows)

    # ---- ključ (INTERNO)
    key1 = p1[["item_id", "group_id", "concept_id", "model", "condition", "sub"]].rename(columns={"item_id": "id"})
    key1["part"] = 1
    key2 = p2[["pair_id", "group_id", "concept_id", "model", "condition", "sub"]].rename(columns={"pair_id": "id"})
    key2["part"] = 2
    key = pd.concat([key1, key2])[["part", "id", "group_id", "concept_id", "model", "condition", "sub"]]
    key.to_csv(out / "annotation_key_INTERNAL.csv", index=False)

    # ---- PDF-ovi
    cover1 = [s for s in CODEBOOK if not s[0].startswith("Part 2")]
    cover2 = [CODEBOOK[0], CODEBOOK[1], CODEBOOK[5], CODEBOOK[6]]
    build_pdf(out / "Part1_booklet_BLIND.pdf", "AccessTrap-CS annotation - Part 1",
              "Rate each response independently. Enter ratings in the Excel workbook (sheet Part1).",
              cover1, [dict(id=r.item_id, domain=r.domain, concept=r.concept, ground_truth=r.ground_truth,
                            response=r.response) for r in p1.itertuples()], "single",
              "AccessTrap-CS | Part 1 | confidential, do not distribute")
    build_pdf(out / "Part2_booklet_OER.pdf", "AccessTrap-CS annotation - Part 2 (OER)",
              "Open this booklet only after you have completed Part 1. Enter ratings in sheet Part2_OER.",
              cover2, blocks, "pair",
              "AccessTrap-CS | Part 2 | confidential, do not distribute")

    # ---- Excel po anotatoru + zip paketi
    for k in (1, 2):
        xp = out / f"entry_rater{k}.xlsx"
        build_xlsx(xp, p1.item_id.tolist(), p2.pair_id.tolist(), f"rater {k}")
        with zipfile.ZipFile(out / f"rater{k}_package.zip", "w", zipfile.ZIP_DEFLATED) as z:
            z.write(out / "Part1_booklet_BLIND.pdf", "Part1_booklet.pdf")
            z.write(out / "Part2_booklet_OER.pdf", "Part2_booklet.pdf")
            z.write(xp, f"entry_rater{k}.xlsx")
    build_tex(out / "codebook_appendix.tex")

    print(f"Uzorak: {len(sel)} grupa, {len(p1)} odgovora (Dio 1), {len(p2)} pojednostavljenih odgovora u {len(blocks)} blokova (Dio 2)")
    print("Po subpolju:", sel["sub"].value_counts().to_dict())
    print("Po modelu  :", sel["model"].value_counts().to_dict())
    print(f"Izlaz u {out}/ :  rater1_package.zip, rater2_package.zip (šalji anotatorima),")
    print("   annotation_key_INTERNAL.csv (NE šalji), codebook_appendix.tex (Appendix B)")


# --------------------------------------------------------------------------- #
# COLLECT
# --------------------------------------------------------------------------- #
def _read(path, key):
    p1 = pd.read_excel(path, sheet_name="Part1")
    p2 = pd.read_excel(path, sheet_name="Part2_OER")
    p1.columns = ["id", "SM", "DA", "accuracy", "comment"][:len(p1.columns)]
    p2.columns = ["id", "OER", "comment"][:len(p2.columns)]
    problems = []
    for df, cols, rng_ in [(p1, ["SM", "DA"], {0, 1}), (p1, ["accuracy"], {1, 2, 3, 4, 5}), (p2, ["OER"], {0, 1})]:
        for col in cols:
            v = pd.to_numeric(df[col], errors="coerce")
            bad = df.loc[v.isna() | ~v.isin(rng_), "id"].tolist()
            if bad:
                problems.append(f"{col}: nevaljane/prazne vrijednosti u {bad[:8]}{'...' if len(bad) > 8 else ''}")
            df[col] = v
    return p1, p2, problems


def cmd_collect(a):
    out = Path(a.outdir)
    key = pd.read_csv(a.key)
    k1 = key[key.part == 1].set_index("id")
    k2 = key[key.part == 2].set_index("id")
    longs = []
    for label, path in [("h1", a.rater1), ("h2", a.rater2)]:
        p1, p2, problems = _read(path, key)
        print(f"[{label}] {path}: Dio1 {len(p1)} redaka, Dio2 {len(p2)} redaka")
        for p in problems:
            print("   UPOZORENJE:", p)
        d1 = p1.merge(k1.reset_index(), on="id", how="left")
        d1["OER"] = np.nan
        d2 = p2.merge(k2.reset_index(), on="id", how="left")[["concept_id", "model", "condition", "OER"]]
        d1 = d1.drop(columns=["OER"]).merge(d2, on=["concept_id", "model", "condition"], how="left")
        d1["rater"] = label
        longs.append(d1[["rater", "concept_id", "model", "condition", "SM", "DA", "OER", "accuracy", "comment"]])
    L = pd.concat(longs, ignore_index=True)
    L.to_csv(out / "annotations_long.csv", index=False)

    keys = ["concept_id", "model", "condition"]
    w = L.pivot_table(index=keys, columns="rater", values=["SM", "DA", "OER"], aggfunc="first")
    print("\nSuglasnost anotatora (sirova % i Cohenov kappa):")
    rows = []
    for y in ["SM", "DA", "OER"]:
        x = w[y][["h1", "h2"]].dropna()
        po = (x.h1 == x.h2).mean()
        pe = sum((x.h1 == v).mean() * (x.h2 == v).mean() for v in (0, 1))
        kap = (po - pe) / (1 - pe) if pe < 1 else np.nan
        print(f"  {y:<4} n={len(x)}  agreement={100*po:.1f}%  kappa={kap:.2f}")
        for idx, r in x[x.h1 != x.h2].iterrows():
            rows.append(dict(indicator=y, concept_id=idx[0], model=idx[1], condition=idx[2], h1=int(r.h1), h2=int(r.h2)))
    adj = pd.DataFrame(rows)
    if len(adj):
        cm = L.set_index(keys + ["rater"])["comment"]
        adj["comment_h1"] = [cm.get((r.concept_id, r.model, r.condition, "h1"), "") for r in adj.itertuples()]
        adj["comment_h2"] = [cm.get((r.concept_id, r.model, r.condition, "h2"), "") for r in adj.itertuples()]
        adj["consensus"] = ""
    adj.to_csv(out / "adjudication_template.csv", index=False)
    print(f"\nNesuglasica za razrješavanje: {len(adj)} -> adjudication_template.csv (popuni stupac 'consensus' 0/1)")
    print("annotations_long.csv spreman. Sljedeći korak: finalize.")


# --------------------------------------------------------------------------- #
# FINALIZE
# --------------------------------------------------------------------------- #
def cmd_finalize(a):
    out = Path(a.outdir)
    L = pd.read_csv(a.long)
    adj = pd.read_csv(a.adjudication) if a.adjudication and Path(a.adjudication).exists() else pd.DataFrame()
    keys = ["concept_id", "model", "condition"]
    cons = {}
    for y in ["SM", "DA", "OER"]:
        w = L.pivot_table(index=keys, columns="rater", values=y, aggfunc="first")
        w = w.dropna(subset=["h1", "h2"])
        for idx, r in w.iterrows():
            if r.h1 == r.h2:
                cons[(idx, y)] = int(r.h1)
    if len(adj):
        adj = adj.dropna(subset=["consensus"])
        for r in adj.itertuples():
            cons[((r.concept_id, r.model, r.condition), r.indicator)] = int(r.consensus)
    rows = {}
    for (idx, y), v in cons.items():
        rows.setdefault(idx, {})[y] = v
    C = pd.DataFrame([dict(rater="consensus", concept_id=i[0], model=i[1], condition=i[2], **v)
                      for i, v in rows.items()])
    # provjera da su sve nesuglasice razriješene
    miss = []
    for y in ["SM", "DA", "OER"]:
        w = L.pivot_table(index=keys, columns="rater", values=y, aggfunc="first").dropna(subset=["h1", "h2"])
        for idx, r in w[w.h1 != w.h2].iterrows():
            if (idx, y) not in cons:
                miss.append((y,) + idx)
    if miss:
        print(f"UPOZORENJE: {len(miss)} nesuglasica nema 'consensus' (npr. {miss[:3]}); te ćelije izostaju iz konsenzusa.")
    F = pd.concat([L.drop(columns=["comment"], errors="ignore"), C], ignore_index=True)
    F.to_csv(out / "annotations_final.csv", index=False)
    print(f"annotations_final.csv: {len(F)} redaka (h1, h2, consensus). Koristi: "
          f"python accesstrap_stats.py --annotations {out/'annotations_final.csv'}")


def main():
    ap = argparse.ArgumentParser()
    sp = ap.add_subparsers(dest="cmd", required=True)
    b = sp.add_parser("build")
    b.add_argument("--corpus", default="corpus.csv")
    b.add_argument("--evals", default="evaluations.csv")
    b.add_argument("--outdir", default="annotation")
    b.add_argument("--n-groups", type=int, default=30)
    b.add_argument("--seed", type=int, default=42)
    c = sp.add_parser("collect")
    c.add_argument("--rater1", required=True)
    c.add_argument("--rater2", required=True)
    c.add_argument("--key", required=True)
    c.add_argument("--outdir", default="annotation")
    f = sp.add_parser("finalize")
    f.add_argument("--long", required=True)
    f.add_argument("--adjudication", default=None)
    f.add_argument("--outdir", default="annotation")
    a = ap.parse_args()
    {"build": cmd_build, "collect": cmd_collect, "finalize": cmd_finalize}[a.cmd](a)


if __name__ == "__main__":
    main()
